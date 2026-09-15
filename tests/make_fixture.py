#!/usr/bin/env python3
"""生成 report_tools.py 的自测夹具：模板 + 3 个 SN 目录的 XML。"""
import json
import os
import sys

from openpyxl import Workbook

OUT = sys.argv[1] if len(sys.argv) > 1 else "/tmp/report-fixture"
os.makedirs(OUT, exist_ok=True)
dataset = os.path.join(OUT, "TD28")
os.makedirs(dataset, exist_ok=True)

CASES = [
    "rt@ver_gain-e1lg_g5_inp18_sgt0_mc",
    "rt@ver_gain-e1lg_g5_inp19_sgt0_mc",
    "rt@ver_gain-e1lg_g5_inp20_sgt0_mc",
    "rt@ver_gain-e1lg_g5_inp21_sgt0_mc",  # 模板里有、XML 里没有 -> Warning
]
VALUES = {
    "ge": "0.022705200000", "te": "0.000000000000",
    "nf": "6.607658430000", "gr": "0.000000000000",
}

# ---- template ----
wb = Workbook()
ws = wb.active
ws.title = "E1&2"
ws.merge_cells("A1:P1")
ws["A1"] = "FIR Report"
headers = {"A3": "No", "B3": "Test Item", "F3": "GE", "H3": "GR", "J3": "TE", "L3": "NF", "P3": "Case"}
for ref, text in headers.items():
    ws[ref] = text
for i, case in enumerate(CASES):
    row = 5 + i
    ws[f"A{row}"] = i + 1
    ws[f"P{row}"] = case
ws["G5"] = "=F5*2"          # 命中行上的目标列公式 -> Write Safety Gate 必须阻断
ws["G8"] = "=F8*2"          # 非目标列的公式，必须原样保留
ws["M9"] = "keep-me"        # 非目标单元格固定文本
ws2 = wb.create_sheet("Summary")
ws2["A1"] = "summary sheet"
ws2["P2"] = "no-case-here"
wb.save(os.path.join(OUT, "template.xlsx"))

# ---- dataset ----
for idx, sn in enumerate(["K7893981", "K7893982", "K7893983"]):
    sn_dir = os.path.join(dataset, sn)
    os.makedirs(sn_dir, exist_ok=True)
    tasks = []
    for ci, case in enumerate(CASES[:3]):
        items = "".join(
            f'    <item name="{name}" value="{float(value) + idx + ci:.12f}"/>\n'
            for name, value in VALUES.items()
        )
        tasks.append(f'  <task name="{case}">\n{items}  </task>\n')
    # 最后一份报告模拟现场常见问题：内容实际 GBK，XML 声明却错误写为 UTF-8。
    description = ' description="中文报告"' if idx == 2 else ""
    xml = '<?xml version="1.0" encoding="UTF-8"?>\n<test_report sn="{}"{}>\n{}</test_report>\n'.format(sn, description, "".join(tasks))
    report_path = os.path.join(sn_dir, f"{sn}_test_report.xml")
    if idx == 2:
        with open(report_path, "wb") as fh:
            fh.write(xml.encode("gbk"))
    else:
        with open(report_path, "w", encoding="utf-8") as fh:
            fh.write(xml)

with open(os.path.join(OUT, "plan.json"), "w", encoding="utf-8") as fh:
    fh.write("""{
  "plan_version": 2,
  "output": {"mode": "one_file_per_sn", "sn_source": "parent_directory_name", "filename_pattern": "{sn}_FIR.xlsx"},
  "source_file": {"pattern": "*_test_report*.xml", "multiple_match_policy": "latest_mtime"},
  "source_records": {
    "record_selector": ".//task",
    "key_selector": "@name",
    "field_storage": {"item_selector": "./item", "field_name_selector": "@name", "field_value_selector": "@value"}
  },
  "template": {
    "sheet_scope": "all",
    "join": {"template_key_column": "P", "source_key": "task.@name", "match_mode": "normalized_exact"},
    "mappings": [
      {"source_field": "ge", "target_column": "F"},
      {"source_field": "gr", "target_column": "H"},
      {"source_field": "te", "target_column": "J"},
      {"source_field": "nf", "target_column": "L"}
    ]
  },
  "normalization": {"trim_whitespace": true, "case_sensitive": true, "rules": []},
  "missing_value_policy": "leave_blank_and_report",
  "missing_case_policy": "warning",
  "duplicate_source_key_policy": "block"
}
""")

base = open(os.path.join(OUT, "plan.json"), encoding="utf-8").read()
with open(os.path.join(OUT, "plan-bad.json"), "w", encoding="utf-8") as fh:
    fh.write(base.replace('".//task"', '".//TestCase"').replace('"@name"', '"@caseName"', 1))


def write_variant(name, mutate):
    plan = json.loads(base)
    mutate(plan)
    with open(os.path.join(OUT, name), "w", encoding="utf-8") as fh:
        json.dump(plan, fh, ensure_ascii=False, indent=2)


# 目标列落在命中行的公式上 -> Write Safety Gate 必须阻断
write_variant("plan-formula.json",
              lambda p: p["template"]["mappings"].append({"source_field": "ge", "target_column": "G"}))
# 不支持的 XPath 语法 -> PLAN_SELECTOR_UNSUPPORTED
write_variant("plan-predicate.json",
              lambda p: p["source_records"].update({"record_selector": ".//task[@name]"}))

print("fixture ready:", OUT)
